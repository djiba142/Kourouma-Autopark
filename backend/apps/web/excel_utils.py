import os
from django.conf import settings
from openpyxl.drawing.image import Image

def add_logo_to_sheet(ws, sheet_title="RAPPORT"):
    """
    Ajoute le logo BANGALY-AUTOPARK-BSG en haut d'une feuille Excel
    et insère quelques lignes pour l'en-tête.
    """
    # 1. Insérer 5 lignes en haut pour l'espace logo/titre
    ws.insert_rows(1, 5)
    
    # 2. Chemin du logo
    logo_path = os.path.join(settings.BASE_DIR, 'static', 'web', 'img', 'logo_bangaly.png')
    
    if os.path.exists(logo_path):
        try:
            img = Image(logo_path)
            # Redimensionner le logo (en pixels, environ 80x80)
            img.width = 80
            img.height = 80
            # Ajouter à la cellule A1 (qui est maintenant vide grâce à l'insertion)
            ws.add_image(img, 'A1')
        except Exception:
            pass # Si Pillow n'est pas bien configuré ou image corrompue
            
    # 3. Ajouter le titre de l'entreprise à côté du logo
    from openpyxl.styles import Font
    ws['C2'] = "BANGALY-AUTOPARK-BSG"
    ws['C2'].font = Font(bold=True, size=16, color="1A3A5C")
    ws['C3'] = sheet_title
    ws['C3'].font = Font(bold=True, size=12, color="666666")
