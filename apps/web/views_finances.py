"""
web/views_finances.py — MODULE FINANCES COMPLET
Admin voit dépenses ENTREPRISE + PERSONNEL + journal complet
Employé voit uniquement dépenses ENTREPRISE
"""
import logging
from datetime import date, timedelta
from django.shortcuts import render, redirect
from django.contrib import messages
from django.db.models import Sum, Q
from django.http import HttpResponse, JsonResponse
from django.utils import timezone

from web.middleware import login_required_web, employe_or_admin, admin_required
from finances.models import JournalFinancier, Depense, CategorieDepense

logger = logging.getLogger(__name__)


def _get_periode_bounds(periode):
    today = date.today()
    if periode == 'jour':
        return today, today
    if periode == 'semaine':
        debut = today - timedelta(days=today.weekday())
        return debut, debut + timedelta(days=6)
    if periode == 'trimestre':
        mois_debut = ((today.month - 1) // 3) * 3 + 1
        debut = today.replace(month=mois_debut, day=1)
        # Fin trimestre = 3 mois après debut - 1 jour
        m = mois_debut + 3
        y = today.year + (m - 1) // 12
        m = ((m - 1) % 12) + 1
        fin = date(y, m, 1) - timedelta(days=1)
        return debut, fin
    if periode == 'annee':
        return today.replace(month=1, day=1), today.replace(month=12, day=31)
    # Défaut : mois courant
    m = today.month
    y = today.year
    if m == 12:
        fin = date(y, 12, 31)
    else:
        fin = date(y, m + 1, 1) - timedelta(days=1)
    return today.replace(day=1), fin


def _journal_qs(debut, fin):
    return JournalFinancier.objects.filter(
        created_at__date__gte=debut, created_at__date__lte=fin
    )


def _depenses_qs(debut, fin, user):
    qs = Depense.objects.filter(
        date_depense__date__gte=debut, date_depense__date__lte=fin
    )
    if user.role != 'ADMIN':
        qs = qs.filter(type_depense='ENTREPRISE')
    return qs


# 1 — RAPPORT FINANCIER PRINCIPAL
@employe_or_admin
def rapport_financier(request):
    periode = request.GET.get('periode', 'mois')
    if periode not in ('jour', 'semaine', 'mois', 'trimestre', 'annee'):
        periode = 'mois'

    debut, fin = _get_periode_bounds(periode)
    is_admin   = (request.user.role == 'ADMIN')

    journal_qs = _journal_qs(debut, fin)
    totaux = journal_qs.aggregate(
        entrees=Sum('montant_gnf', filter=Q(type_flux='ENTREE')),
        sorties=Sum('montant_gnf', filter=Q(type_flux='SORTIE')),
    )
    entrees = float(totaux['entrees'] or 0)
    sorties = float(totaux['sorties'] or 0)
    profit  = entrees - sorties

    # Solde caisse global (toutes périodes)
    solde_all = JournalFinancier.objects.aggregate(
        e=Sum('montant_gnf', filter=Q(type_flux='ENTREE')),
        s=Sum('montant_gnf', filter=Q(type_flux='SORTIE')),
    )
    solde_caisse = float(solde_all['e'] or 0) - float(solde_all['s'] or 0)

    dep_qs = _depenses_qs(debut, fin, request.user)
    dep_entreprise = float(dep_qs.filter(type_depense='ENTREPRISE').aggregate(t=Sum('montant_gnf'))['t'] or 0)
    dep_personnel  = float(dep_qs.filter(type_depense='PERSONNEL').aggregate(t=Sum('montant_gnf'))['t'] or 0) if is_admin else 0

    dep_par_categorie = list(
        dep_qs.values('categorie__nom').annotate(total=Sum('montant_gnf')).order_by('-total')[:6]
    )

    # Graphique 7 derniers jours
    labels_chart, data_ventes, data_depenses = [], [], []
    for i in range(6, -1, -1):
        d = date.today() - timedelta(days=i)
        labels_chart.append(d.strftime('%d %b'))
        data_ventes.append(float(JournalFinancier.objects.filter(
            created_at__date=d, type_flux='ENTREE',
            type_operation__in=['VENTE', 'PAIEMENT']
        ).aggregate(t=Sum('montant_gnf'))['t'] or 0))
        data_depenses.append(float(JournalFinancier.objects.filter(
            created_at__date=d, type_flux='SORTIE'
        ).aggregate(t=Sum('montant_gnf'))['t'] or 0))

    # Budget par catégorie
    categories_budget = []
    for cat in CategorieDepense.objects.filter(budget_mensuel__gt=0):
        consomme = float(dep_qs.filter(categorie=cat).aggregate(t=Sum('montant_gnf'))['t'] or 0)
        pct = min(int(consomme / float(cat.budget_mensuel) * 100), 100) if cat.budget_mensuel else 0
        
        # Détermination de la classe CSS dans le backend
        if pct >= 100: cat_class = 'bg-danger'
        elif pct >= 80: cat_class = 'bg-warning'
        else: cat_class = 'bg-success'
        
        categories_budget.append({
            'nom': cat.nom, 'budget': float(cat.budget_mensuel),
            'consomme': consomme, 'pct': pct, 'alerte': pct >= 80,
            'cat_class': cat_class,
        })

    # Données JSON pour Chart.js
    import json
    labels_json  = json.dumps(labels_chart)
    ventes_json  = json.dumps(data_ventes)
    depenses_json = json.dumps(data_depenses)

    return render(request, 'web/finances/rapport.html', {
        'title': 'Rapport Financier', 'periode': periode,
        'debut': debut, 'fin': fin,
        'entrees': entrees, 'sorties': sorties, 'profit': profit,
        'solde_caisse': solde_caisse,
        'dep_entreprise': dep_entreprise, 'dep_personnel': dep_personnel,
        'dep_total': dep_entreprise + dep_personnel,
        'dep_par_categorie': dep_par_categorie,
        'categories_budget': categories_budget,
        'labels_json': labels_json,
        'ventes_json': ventes_json,
        'depenses_json': depenses_json,
        'journal_recent': journal_qs.select_related('utilisateur').order_by('-created_at')[:20],
        'is_admin': is_admin,
    })


# 2 — LISTE DÉPENSES filtrée par rôle
@employe_or_admin
def depenses_liste(request):
    periode  = request.GET.get('periode', 'mois')
    type_dep = request.GET.get('type', '')
    cat_id   = request.GET.get('cat', '')
    debut, fin = _get_periode_bounds(periode)
    qs = _depenses_qs(debut, fin, request.user).select_related('categorie', 'cree_par')
    if type_dep: qs = qs.filter(type_depense=type_dep)
    if cat_id:   qs = qs.filter(categorie_id=cat_id)

    is_admin = (request.user.role == 'ADMIN')
    total_ent = float(qs.filter(type_depense='ENTREPRISE').aggregate(t=Sum('montant_gnf'))['t'] or 0)
    total_per = float(qs.filter(type_depense='PERSONNEL').aggregate(t=Sum('montant_gnf'))['t'] or 0) if is_admin else 0

    return render(request, 'web/finances/depenses.html', {
        'title': 'Dépenses', 'depenses': qs.order_by('-date_depense')[:200],
        'total': total_ent + total_per, 'total_entreprise': total_ent, 'total_personnel': total_per,
        'categories': CategorieDepense.objects.all().order_by('nom'),
        'periode': periode, 'type_dep': type_dep, 'cat_id': cat_id, 'is_admin': is_admin,
    })


# 3 — NOUVELLE DÉPENSE
@employe_or_admin
def depense_nouvelle(request):
    is_admin = (request.user.role == 'ADMIN')
    if request.method == 'GET':
        return render(request, 'web/finances/depense_form.html', {
            'title': 'Nouvelle Dépense',
            'categories': CategorieDepense.objects.all().order_by('nom'),
            'types': Depense.TYPES, 'devises': Depense.DEVISES, 'is_admin': is_admin,
        })

    montant     = request.POST.get('montant', '').strip()
    description = request.POST.get('description', '').strip()
    cat_nom     = request.POST.get('categorie_nom', '').strip()
    type_dep    = request.POST.get('type_depense', 'ENTREPRISE')
    devise      = request.POST.get('devise', 'GNF')
    taux        = request.POST.get('taux', '1')
    preuve      = request.FILES.get('preuve')

    errors = []
    try:
        montant_f = float(montant)
        if montant_f <= 0: errors.append('Montant invalide.')
    except (ValueError, TypeError):
        errors.append('Montant invalide.'); montant_f = 0

    if not description: errors.append('Description obligatoire.')
    if not cat_nom:      errors.append('Catégorie obligatoire.')
    if type_dep == 'PERSONNEL' and not is_admin:
        errors.append('Seul un administrateur peut enregistrer une dépense personnelle.')
        type_dep = 'ENTREPRISE'

    if errors:
        for e in errors: messages.error(request, e)
        return render(request, 'web/finances/depense_form.html', {
            'title': 'Nouvelle Dépense',
            'categories': CategorieDepense.objects.all().order_by('nom'),
            'types': Depense.TYPES, 'devises': Depense.DEVISES, 'is_admin': is_admin,
        })

    try: taux_f = float(taux)
    except ValueError: taux_f = 1.0

    # Créer la catégorie si elle n'existe pas
    cat_obj, _ = CategorieDepense.objects.get_or_create(nom=cat_nom)

    Depense.objects.create(
        montant=montant_f, description=description, categorie=cat_obj,
        type_depense=type_dep, devise=devise, taux=taux_f,
        preuve=preuve, cree_par=request.user,
    )
    label = 'Personnelle' if type_dep == 'PERSONNEL' else 'Professionnelle'
    messages.success(request, f'Dépense {label} de {montant_f:,.0f} GNF enregistrée.')
    return redirect('web:depenses_liste')


# 4 — EXPORT EXCEL (Admin only)
@admin_required
def export_journal_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    periode = request.GET.get('periode', 'mois')
    debut, fin = _get_periode_bounds(periode)

    wb  = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = 'Journal'
    hf  = PatternFill('solid', fgColor='1A3A5C')
    hfont = Font(bold=True, color='FFFFFF')
    ctr = Alignment(horizontal='center')

    cols = ['N°','Date','Type','Flux','Référence','Description','Montant GNF','Par']
    widths = [20,18,16,10,22,40,18,20]
    for c,(h,w) in enumerate(zip(cols,widths),1):
        cell = ws1.cell(1,c,h); cell.font=hfont; cell.fill=hf; cell.alignment=ctr
        ws1.column_dimensions[cell.column_letter].width = w

    gf = Font(bold=True,color='1E7E34'); rf = Font(bold=True,color='C0392B')
    for r,j in enumerate(JournalFinancier.objects.filter(
        created_at__date__gte=debut,created_at__date__lte=fin
    ).select_related('utilisateur').order_by('-created_at'),2):
        ws1.cell(r,1,j.numero); ws1.cell(r,2,j.created_at.strftime('%d/%m/%Y %H:%M'))
        ws1.cell(r,3,j.type_operation); ws1.cell(r,4,j.type_flux)
        ws1.cell(r,5,j.reference_doc); ws1.cell(r,6,j.description[:80])
        mc=ws1.cell(r,7,float(j.montant_gnf)); mc.font=gf if j.type_flux=='ENTREE' else rf
        ws1.cell(r,8,j.utilisateur.nom if j.utilisateur else '—')

    ws2 = wb.create_sheet('Dépenses')
    cols2=['N°','Description','Catégorie','Type','Montant GNF','Date','Par']
    for c,h in enumerate(cols2,1):
        cell=ws2.cell(1,c,h); cell.font=hfont; cell.fill=hf
    for r,d in enumerate(Depense.objects.filter(
        date_depense__date__gte=debut,date_depense__date__lte=fin
    ).select_related('categorie','cree_par').order_by('-date_depense'),2):
        ws2.cell(r,1,d.numero or '—'); ws2.cell(r,2,d.description[:80])
        ws2.cell(r,3,d.categorie.nom if d.categorie else '—')
        ws2.cell(r,4,d.type_depense); ws2.cell(r,5,float(d.montant_gnf))
        ws2.cell(r,6,d.date_depense.strftime('%d/%m/%Y'))
        ws2.cell(r,7,d.cree_par.nom if d.cree_par else '—')

    from web.excel_utils import add_logo_to_sheet
    add_logo_to_sheet(ws1, f"Journal Financier ({periode})")
    add_logo_to_sheet(ws2, "Liste des Dépenses")

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="BANGALY_Rapport_{periode}_{debut.strftime("%Y%m%d")}.xlsx"'
    wb.save(resp)
    return resp


# 5 — SOLDE CAISSE JSON (HTMX)
@employe_or_admin
def solde_caisse_json(request):
    s = JournalFinancier.objects.aggregate(
        e=Sum('montant_gnf',filter=Q(type_flux='ENTREE')),
        s=Sum('montant_gnf',filter=Q(type_flux='SORTIE')),
    )
    total = float(s['e'] or 0) - float(s['s'] or 0)
    return JsonResponse({'solde': total, 'solde_fmt': f"{total:,.0f} GNF"})


# Alias pour compatibilité avec urls.py existant
finances_liste = rapport_financier
